import io
from docx import Document
def parse(data):
 d=Document(io.BytesIO(data));parts=[p.text for p in d.paragraphs if p.text.strip()]
 for t in d.tables:
  parts += [' | '.join(c.text.strip() for c in row.cells) for row in t.rows]
 return '\n'.join(parts),1
