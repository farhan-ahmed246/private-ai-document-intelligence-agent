import io
from openpyxl import load_workbook
def parse(data):
 w=load_workbook(io.BytesIO(data),read_only=True,data_only=True);parts=[]
 for s in w.worksheets:
  parts.append('[SHEET '+s.title+']')
  for row in s.iter_rows(values_only=True):
   vals=['' if x is None else str(x) for x in row]
   if any(vals):parts.append(' | '.join(vals))
 return '\n'.join(parts),len(w.worksheets)
